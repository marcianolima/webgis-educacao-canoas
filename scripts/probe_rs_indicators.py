import io,time,urllib.request
import openpyxl
RESOURCES={
"complexidade":"https://dados.rs.gov.br/dataset/8040e08b-144a-43d2-a6a7-5514a638b4f0/resource/ae5e3e52-7bed-4c48-a4d4-07d138ebb8ab/download/dados_complexidade_gestao.xlsx",
"inse":"https://dados.rs.gov.br/dataset/8040e08b-144a-43d2-a6a7-5514a638b4f0/resource/a19c1327-63d4-4c13-b3e8-b5b48922312a/download/dados_inse.xlsx",
"afd":"https://dados.rs.gov.br/dataset/8040e08b-144a-43d2-a6a7-5514a638b4f0/resource/6bd1185a-8fa6-49cc-a980-56360e302e6b/download/dados_afd.xlsx",
"rendimento":"https://dados.rs.gov.br/dataset/8040e08b-144a-43d2-a6a7-5514a638b4f0/resource/0b3ea083-b442-4e4a-b679-64cb7cc6a76f/download/dados_fluxo_rendimento.xlsx",
"distorcao":"https://dados.rs.gov.br/dataset/8040e08b-144a-43d2-a6a7-5514a638b4f0/resource/939cd0e3-0e2f-4064-a315-4dcc849cf8ec/download/dados_tdi.xlsx"
}
for name,url in RESOURCES.items():
 print("RESOURCE",name,flush=True)
 payload=None
 for attempt in range(3):
  try:
   req=urllib.request.Request(url,headers={"User-Agent":"Mozilla/5.0"})
   with urllib.request.urlopen(req,timeout=45) as r:payload=r.read(12_000_000)
   break
  except Exception as e:
   print("RETRY",name,str(e),flush=True)
   time.sleep(4)
 if payload is None:continue
 print("BYTES",len(payload),flush=True)
 wb=openpyxl.load_workbook(io.BytesIO(payload),read_only=True,data_only=True)
 for ws in wb:
  if ws.title.lower() in ("sobre","dicionario"):continue
  it=ws.values
  header=list(next(it,[]))
  print("SHEET",ws.title,"NROWS",ws.max_row,"COL",ws.max_column,"HEAD",[str(x)[:45] for x in header[:20]],flush=True)
  sample=[]
  for i,row in enumerate(it):
   if "4304606" in [str(x).removesuffix(".0") for x in row[:5]] or "Canoas" in row[:5]:
    sample.append([str(x)[:45] for x in row[:18]])
   if i>25000:break
  if sample:
   print("CANOAS",len(sample),"FIRST",sample[:2],"LAST",sample[-2:],flush=True)
 wb.close()
